"""Informalidad laboral (DANE, GEIH - empleo informal y seguridad social).

Salidas
-------
informalidad.csv        proporcion de ocupados informales (%), trimestre movil, desde 2021:
                        fecha (ultimo mes del trimestre movil), periodo, nacional, ciudades_13, ciudades_23
informalidad_ramas.csv  total nacional por rama de actividad: fecha, rama, ocupados, informales, tasa (%)
                        (miles de personas; tasa = informales / ocupados)

Reglas
------
* La serie comienza en 2021 (medicion con el marco muestral 2018 y la definicion vigente del DANE).
* Cada columna es un trimestre movil consecutivo; se valida que el mes final de cada rotulo
  coincida con la secuencia (los rotulos del DANE no siempre traen el ano).
* Si la descarga falla se conservan los archivos anteriores.

Uso: python -m colombiamacro.fuentes.informalidad
"""

from __future__ import annotations

import io
import re
import sys
import unicodedata
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from colombiamacro.config import data

OUT = data("informalidad.csv")
OUT_RAMAS = data("informalidad_ramas.csv")
OUT_CIUDADES = data("informalidad_ciudades.csv")
# Las 13 ciudades y areas metropolitanas (A.M.) de la medicion tradicional del DANE
CIUDADES_13 = ["Bogotá D.C.", "Medellín A.M.", "Cali A.M.", "Barranquilla A.M.", "Bucaramanga A.M.",
               "Manizales A.M.", "Pasto", "Pereira A.M.", "Cúcuta A.M.", "Ibagué", "Montería",
               "Cartagena", "Villavicencio"]
PAGINA = ("https://www.dane.gov.co/index.php/estadisticas-por-tema/mercado-laboral/"
          "empleo-informal-y-seguridad-social")
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124.0 Safari/537.36"}
MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
         "sep": 9, "oct": 10, "nov": 11, "dic": 12}
DOMINIOS = {"total nacional": "nacional", "13 ciudades y a.m.": "ciudades_13", "23 ciudades y a.m.": "ciudades_23"}
RAMAS_CORTAS = {
    "agricultura": "Agro y pesca", "suministro de electricidad": "Electricidad, gas y agua",
    "industrias manufactureras": "Industria", "construccion": "Construcción",
    "comercio": "Comercio", "alojamiento": "Alojamiento y comida", "transporte": "Transporte",
    "informacion": "Información y comunicaciones", "actividades financieras": "Finanzas y seguros",
    "actividades inmobiliarias": "Inmobiliarias", "actividades profesionales": "Servicios profesionales",
    "administracion publica": "Gobierno, educación y salud", "actividades artisticas": "Arte, hogares y otros",
}


def _norm(txt) -> str:
    t = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", t).strip()


def rama_corta(nombre: str) -> str | None:
    n = _norm(nombre)
    for clave, corto in RAMAS_CORTAS.items():
        if n.startswith(clave):
            return corto
    return None


def fechas_columnas(anos: tuple, rotulos: tuple) -> dict[int, pd.Timestamp]:
    """Columna -> ultimo mes del trimestre movil. Secuencia mensual validada con cada rotulo."""
    cols = [c for c in range(1, len(rotulos)) if rotulos[c] and re.search(r"[A-Za-z]{3}", str(rotulos[c]))]
    if not cols:
        raise ValueError("No hay columnas de trimestre movil")
    primer_ano = next((int(a) for a in anos[1:] if isinstance(a, (int, float)) or str(a).strip().isdigit()), None)
    if primer_ano is None:
        raise ValueError("No se encontro el primer ano de la serie")
    out = {}
    for k, c in enumerate(cols):
        fin = re.findall(r"([A-Za-z]{3})", _norm(rotulos[c]))[-1]
        if k == 0:
            actual = pd.Timestamp(primer_ano, MESES[fin], 1)
        else:
            actual = actual + pd.DateOffset(months=1)
            if actual.month != MESES[fin]:
                raise ValueError(f"Secuencia de trimestres moviles rota en '{rotulos[c]}'")
        out[c] = actual
    return out


def _filas(ws) -> list[tuple]:
    return list(ws.values)


def leer_ciudades(wb) -> pd.DataFrame:
    """Proporcion de informales por ciudad (23 ciudades y areas metropolitanas)."""
    hoja = next(n for n in wb.sheetnames if _norm(n).startswith("prop informalidad"))
    filas = _filas(wb[hoja])
    i_enc = next(i for i, r in enumerate(filas) if r[0] and _norm(r[0]).startswith("proporcion de informal") and r[1])
    fechas = fechas_columnas(filas[i_enc], filas[i_enc + 1])
    trece = {_norm(c) for c in CIUDADES_13}
    out = []
    for r in filas[i_enc + 2:]:
        nombre = str(r[0] or "").strip()
        if not nombre or _norm(nombre).startswith(("fuente", "nota")):
            if out:
                break
            continue
        if _norm(nombre) in DOMINIOS:
            continue
        for c, f in fechas.items():
            if isinstance(r[c], (int, float)):
                out.append({"fecha": f, "ciudad": nombre, "tasa": round(float(r[c]), 4),
                            "grupo": "13" if _norm(nombre) in trece else "23"})
    df = pd.DataFrame(out)
    if df["ciudad"].nunique() < 20 or (df.drop_duplicates("ciudad")["grupo"] == "13").sum() != 13:
        raise ValueError("No se encontraron las 23 ciudades (o las 13 principales) en el anexo")
    return df.sort_values(["ciudad", "fecha"]).reset_index(drop=True)


def leer_proporcion(wb) -> pd.DataFrame:
    hoja = next(n for n in wb.sheetnames if _norm(n).startswith("prop informalidad"))
    filas = _filas(wb[hoja])
    i_enc = next(i for i, r in enumerate(filas) if r[0] and _norm(r[0]).startswith("proporcion de informal") and r[1])
    fechas = fechas_columnas(filas[i_enc], filas[i_enc + 1])
    series = {}
    for r in filas[i_enc + 2:]:
        clave = DOMINIOS.get(_norm(r[0]))
        if clave:
            series[clave] = {f: float(r[c]) for c, f in fechas.items() if isinstance(r[c], (int, float))}
    if "nacional" not in series:
        raise ValueError("Falta la proporcion de informalidad total nacional")
    df = pd.DataFrame(series).sort_index()
    df.index.name = "fecha"
    df = df.reset_index()
    if not df["nacional"].between(20, 90).all():
        raise ValueError("Informalidad nacional fuera de rango")
    df.insert(1, "periodo", [f"{f.year}-{f.month:02d}" for f in df["fecha"]])
    return df


def leer_ramas(wb) -> pd.DataFrame:
    hoja = next(n for n in wb.sheetnames if _norm(n).startswith("ramas de actividad"))
    filas = _filas(wb[hoja])
    i_tn = next(i for i, r in enumerate(filas) if _norm(r[0]) == "total nacional")
    fechas = fechas_columnas(filas[i_tn + 1], filas[i_tn + 2])
    bloque, datos = None, {"ocupados": {}, "informales": {}}
    for r in filas[i_tn + 3:]:
        n = _norm(r[0])
        if n in ("", "none") and bloque:
            break
        if n.startswith("poblacion ocupada"):
            bloque = "ocupados"
            continue
        if n == "formal":
            bloque = "formal"
            continue
        if n == "informal":
            bloque = "informales"
            continue
        corto = rama_corta(r[0])
        if corto and bloque in datos:
            datos[bloque][corto] = {f: float(r[c]) for c, f in fechas.items() if isinstance(r[c], (int, float))}
    if len(datos["ocupados"]) < 12 or len(datos["informales"]) < 12:
        raise ValueError("Faltan ramas de actividad en el anexo de informalidad")
    filas_out = []
    for rama, serie in datos["ocupados"].items():
        for f, ocup in serie.items():
            inf = datos["informales"][rama].get(f)
            if inf is None or ocup <= 0:
                continue
            filas_out.append({"fecha": f, "rama": rama, "ocupados": round(ocup, 3), "informales": round(inf, 3),
                              "tasa": round(100 * inf / ocup, 3)})
    return pd.DataFrame(filas_out).sort_values(["rama", "fecha"]).reset_index(drop=True)


def leer_anexo(contenido: bytes) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        return leer_proporcion(wb), leer_ramas(wb), leer_ciudades(wb)
    finally:
        wb.close()


def enlace_anexo(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    enlaces = [urljoin(PAGINA, a["href"]) for a in soup.select("a[href]") if "anex-geiheiss" in a["href"].lower()]
    enlaces = [u for u in enlaces if u.lower().endswith(".xlsx")]
    if not enlaces:
        raise ValueError("No se encontro el anexo de informalidad (anex-GEIHEISS)")
    return enlaces[0]


def escribir(df: pd.DataFrame, ruta) -> bool:
    d = df.copy()
    d["fecha"] = pd.to_datetime(d["fecha"]).dt.strftime("%Y-%m-%d")
    texto = d.to_csv(index=False, lineterminator="\n", float_format="%.4f")
    if ruta.exists() and ruta.read_text(encoding="utf-8") == texto:
        return False
    ruta.write_text(texto, encoding="utf-8")
    return True


def main() -> int:
    s = requests.Session()
    s.headers.update(HEADERS)
    r = s.get(PAGINA, timeout=40)
    r.raise_for_status()
    url = enlace_anexo(r.text)
    r = s.get(url, timeout=120)
    r.raise_for_status()
    prop, ramas, ciudades = leer_anexo(r.content)
    escribir(prop, OUT)
    escribir(ramas, OUT_RAMAS)
    escribir(ciudades, OUT_CIUDADES)
    ult = prop.iloc[-1]
    print(f"  informalidad: {len(prop)} trimestres moviles hasta {ult['periodo']} "
          f"(nacional {ult['nacional']:.1f}%); {ramas['rama'].nunique()} ramas — {url}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # se conservan los archivos anteriores
        print(f"  ERROR informalidad: {exc}")
        sys.exit(1)
