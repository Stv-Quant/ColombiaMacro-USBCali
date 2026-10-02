"""PIB por el enfoque del gasto (DANE) y poblacion nacional (DANE, proyecciones oficiales).

Escribe:
  pib_gasto.csv            componentes de la demanda: nivel real original, real desestacionalizado y nominal
  pib_inversion.csv        formacion bruta de capital fijo por tipo de activo (real original y desestacionalizado)
  pib_consumo_hogares.csv  consumo de los hogares por durabilidad y por finalidad COICOP (real original y desestacionalizado)
  poblacion.csv            poblacion total nacional a mitad de ano (DANE, censo 2018, actualizacion 2025)

Como en pib.py, cada publicacion reemplaza toda la serie (el DANE revisa trimestres anteriores).
Uso: python -m colombiamacro.fuentes.demanda [--local carpeta]  (--local procesa anexos ya descargados)
"""

from __future__ import annotations

import argparse
import io
import re
import sys
import time
import unicodedata
from pathlib import Path
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from colombiamacro.config import data
from colombiamacro.fuentes.banrep import write_if_changed

DANE_PIB = ("https://www.dane.gov.co/index.php/estadisticas-por-tema/"
            "cuentas-nacionales/cuentas-nacionales-trimestrales/pib-informacion-tecnica")
DANE_POB = "https://www.dane.gov.co/index.php/estadisticas-por-tema/demografia-y-poblacion/proyecciones-de-poblacion"
PATRONES = {
    "constantes": r"anex-GastoConstantes-[IVX]+trim20\d{2}\.xlsx$",
    "corrientes": r"anex-GastoCorriente-[IVX]+trim20\d{2}\.xlsx$",
    "pob_reciente": r"PPED-AreaNac-2018-20\d{2}\.xlsx$",
    "pob_historica": r"PPED-AreaNac-1950-2017\.xlsx$",
}
QUARTERS = {"I": 1, "II": 4, "III": 7, "IV": 10}
HEADERS = {"User-Agent": "Mozilla/5.0 (ColombiaMacro; academic dashboard)"}

# Cuadro 1/2: codigo SCN -> clave
GASTO = {"P.8": "demanda_interna", "P.3": "consumo_final", "P.311 + P.313 + P.323": "consumo_hogares",
         "P.312 + P.322": "consumo_gobierno", "P.5": "formacion_capital", "P.51": "inversion_fija",
         "P.6": "exportaciones", "P.7": "importaciones", "B.1b": "pib",
         "P.61": "exportaciones_bienes", "P.62": "exportaciones_servicios",
         "P.71": "importaciones_bienes", "P.72": "importaciones_servicios"}
ACTIVOS = {"AN111": "vivienda", "AN112": "otros_edificios", "AN113 + AN114": "maquinaria_equipo",
           "AN115": "recursos_biologicos", "AN117": "propiedad_intelectual"}
DURABILIDAD = {"bienes durables": "durables", "bienes semidurables": "semidurables",
               "bienes no durables": "no_durables", "servicios": "servicios"}
COICOP = {"01": "alimentos", "02": "alcohol_tabaco", "03": "vestuario", "04": "vivienda_servicios",
          "05": "muebles_hogar", "06": "salud", "07": "transporte", "08": "comunicaciones",
          "09": "recreacion", "10": "educacion", "11": "restaurantes_hoteles", "12": "diversos"}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def _codigo(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def enlace(html: str, pagina: str, patron: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    links = list(dict.fromkeys(a["href"] for a in soup.select("a[href]") if re.search(patron, a["href"], re.I)))
    if len(links) != 1:
        raise ValueError(f"DANE: se esperaba 1 enlace para {patron}; hay {len(links)}")
    url = urljoin(pagina, links[0])
    if urlparse(url).hostname not in ("www.dane.gov.co", "dane.gov.co"):
        raise ValueError(f"URL fuera del DANE: {url}")
    return url


def descargar(url: str, s: requests.Session) -> bytes:
    for intento in range(5):
        r = s.get(url, timeout=90, headers=HEADERS)
        if r.status_code == 429:
            time.sleep(15 * (intento + 1))
            continue
        r.raise_for_status()
        return r.content
    raise RuntimeError(f"DANE limita las descargas (429): {url}")


# ------------------------------------------------------------------ lectura de cuadros trimestrales
def bloque_trimestral(rows: list, clave_de) -> dict[str, pd.Series]:
    """Primer bloque de niveles de un cuadro del DANE: fila con 'Concepto' (anos), fila siguiente
    (trimestres I-IV) y filas de datos hasta 'Fuente'. clave_de(col0, col1) -> clave o None."""
    h = next(i for i, r in enumerate(rows) if r and len(r) > 1 and _norm(r[1]) == "concepto")
    anos, trims = rows[h], rows[h + 1]
    cols, ano = {}, None
    for c in range(2, len(trims)):
        m = re.match(r"^(19|20)\d{2}", str(anos[c] if c < len(anos) else "").strip())
        if m:
            ano = int(str(anos[c]).strip()[:4])
        q = str(trims[c] or "").strip()
        if q in QUARTERS and ano:
            cols[c] = pd.Timestamp(ano, QUARTERS[q], 1)
    out = {}
    for r in rows[h + 2:]:
        if r and _norm(r[0]).startswith("fuente"):
            break
        if not r or len(r) < 3:
            continue
        k = clave_de(r[0], r[1])
        if k and k not in out:
            out[k] = pd.Series({f: float(r[c]) for c, f in cols.items() if c < len(r) and isinstance(r[c], (int, float))})
    if len(cols) < 60:
        raise ValueError(f"Historia trimestral corta: {len(cols)} trimestres")
    return out


def _hoja(contenido: bytes, hoja: str) -> list:
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[hoja].values]
    finally:
        wb.close()


def leer_gasto(constantes: bytes, corrientes: bytes) -> pd.DataFrame:
    k = lambda a, b: GASTO.get(_codigo(a))
    orig = {**bloque_trimestral(_hoja(constantes, "Cuadro 1"), k), **bloque_trimestral(_hoja(constantes, "Cuadro 7"), k)}
    sa = {**bloque_trimestral(_hoja(constantes, "Cuadro 2"), k), **bloque_trimestral(_hoja(constantes, "Cuadro 8"), k)}
    nom = bloque_trimestral(_hoja(corrientes, "Cuadro 1"), k)
    faltan = [c for c in ("pib", "consumo_hogares", "consumo_gobierno", "inversion_fija", "exportaciones", "importaciones") if c not in orig or c not in sa or c not in nom]
    if faltan:
        raise ValueError(f"Gasto: faltan componentes {faltan}")
    filas = []
    for comp, s in orig.items():
        for f, v in s.items():
            filas.append({"fecha": f, "componente": comp, "real": v, "real_sa": sa.get(comp, pd.Series(dtype=float)).get(f),
                          "nominal": nom.get(comp, pd.Series(dtype=float)).get(f)})
    df = pd.DataFrame(filas).sort_values(["componente", "fecha"]).reset_index(drop=True)
    return df


def leer_inversion(constantes: bytes) -> pd.DataFrame:
    k = lambda a, b: ACTIVOS.get(_codigo(a))
    orig, sa = bloque_trimestral(_hoja(constantes, "Cuadro 5"), k), bloque_trimestral(_hoja(constantes, "Cuadro 6"), k)
    if set(orig) != set(ACTIVOS.values()):
        raise ValueError(f"Inversion por activo incompleta: {sorted(orig)}")
    return pd.DataFrame([{"fecha": f, "activo": a, "real": v, "real_sa": sa[a].get(f)} for a, s in orig.items() for f, v in s.items()]
                        ).sort_values(["activo", "fecha"]).reset_index(drop=True)


def leer_consumo(constantes: bytes) -> pd.DataFrame:
    def bloques(hoja):
        rows = _hoja(constantes, hoja)
        fin = bloque_trimestral(rows, lambda a, b: COICOP.get(_codigo(a)))
        # la seccion de durabilidad esta en el mismo bloque, sin codigo
        dur = bloque_trimestral(rows, lambda a, b: DURABILIDAD.get(_norm(b)) if not _codigo(a) else None)
        return fin, dur
    fo, do_ = bloques("Cuadro 3")
    fs, ds = bloques("Cuadro 4")
    if len(fo) != 12 or len(do_) != 4:
        raise ValueError("Consumo de hogares: grupos incompletos")
    filas = []
    for tipo, o, s_ in (("finalidad", fo, fs), ("durabilidad", do_, ds)):
        for g, s in o.items():
            for f, v in s.items():
                filas.append({"fecha": f, "tipo": tipo, "grupo": g, "real": v, "real_sa": s_.get(g, pd.Series(dtype=float)).get(f)})
    return pd.DataFrame(filas).sort_values(["tipo", "grupo", "fecha"]).reset_index(drop=True)


def leer_poblacion(reciente: bytes, historica: bytes | None, hasta: int) -> pd.DataFrame:
    """Total nacional por ano. hasta: ultimo ano a publicar (el del ultimo dato del PIB)."""
    def tabla(cont):
        wb = load_workbook(io.BytesIO(cont), read_only=True, data_only=True)
        try:
            out = {}
            for ws in wb.worksheets:
                for r in ws.values:
                    r = list(r)
                    txt = [_norm(x) for x in r]
                    if "total" not in txt or not any("nacional" in x for x in txt):
                        continue
                    anos = [x for x in r if isinstance(x, (int, float)) and 1900 < x < 2200 and float(x).is_integer()]
                    vals = [x for x in r if isinstance(x, (int, float)) and x > 1e6]
                    if anos and vals:
                        out[int(anos[0])] = float(vals[-1])
            return out
        finally:
            wb.close()
    pob = tabla(historica) if historica else {}
    rec = tabla(reciente)
    pob.update(rec)                                    # 2018 en adelante: actualizacion vigente
    if len(rec) < 20:
        raise ValueError("Poblacion: serie reciente incompleta")
    s = pd.Series(pob).sort_index()
    s = s[s.index <= hasta]
    if not s.between(5e6, 1e8).all():
        raise ValueError("Poblacion fuera de rango")
    return pd.DataFrame({"anio": s.index, "poblacion": s.values.astype("int64"),
                         "tipo": ["proyeccion_oficial" if a >= 2018 else "retroproyeccion"
                                  for a in s.index]})


def validar(g: pd.DataFrame) -> None:
    p = g[g["componente"] == "pib"].set_index("fecha")
    if (p["real"].pct_change(4, fill_method=None).abs() > 0.4).any():
        raise ValueError("Gasto: crecimiento del PIB fuera de rango")
    pib_csv = data("pib_colombia.csv")
    if pib_csv.exists():
        ref = pd.read_csv(pib_csv, parse_dates=["fecha"]).set_index("fecha")["pib_real_miles_millones_ref2015"]
        comun = ref.index.intersection(p.index)
        if len(comun) and ((p.loc[comun, "real"] - ref[comun]).abs() / ref[comun]).max() > 0.002:
            raise ValueError("Gasto: el PIB no coincide con el anexo de produccion (vintages distintos)")


def escribir(gasto, inversion, consumo, poblacion) -> None:
    for df, nombre in ((gasto, "pib_gasto.csv"), (inversion, "pib_inversion.csv"), (consumo, "pib_consumo_hogares.csv")):
        x = df.copy()
        x["fecha"] = pd.to_datetime(x["fecha"]).dt.strftime("%Y-%m-%d")
        ch = write_if_changed(x, data(nombre), float_format="%.4f")
        print(f"{nombre}: {len(x)} filas hasta {x['fecha'].max()} ({'actualizado' if ch else 'sin cambios'})")
    ch = write_if_changed(poblacion, data("poblacion.csv"))
    print(f"poblacion.csv: {poblacion['anio'].min()}–{poblacion['anio'].max()} ({'actualizado' if ch else 'sin cambios'})")


def actualizar(local: Path | None = None) -> bool:
    cont = {}
    if local:
        for k, pat in PATRONES.items():
            hits = [p for p in Path(local).iterdir() if re.search(pat, p.name, re.I)]
            if hits:
                cont[k] = sorted(hits)[-1].read_bytes()
    else:
        s = requests.Session()
        for pagina, claves in ((DANE_PIB, ("constantes", "corrientes")), (DANE_POB, ("pob_reciente", "pob_historica"))):
            html = s.get(pagina, timeout=40, headers=HEADERS)
            html.raise_for_status()
            for k in claves:
                try:
                    cont[k] = descargar(enlace(html.text, pagina, PATRONES[k]), s)
                    time.sleep(4)
                except Exception as exc:
                    print(f"  ERROR descarga {k}: {exc}")
    ok = True
    try:
        gasto = leer_gasto(cont["constantes"], cont["corrientes"])
        validar(gasto)
        inversion, consumo = leer_inversion(cont["constantes"]), leer_consumo(cont["constantes"])
        hasta = int(pd.to_datetime(gasto["fecha"]).max().year)
        poblacion = leer_poblacion(cont["pob_reciente"], cont.get("pob_historica"), hasta)
        escribir(gasto, inversion, consumo, poblacion)
    except Exception as exc:
        ok = False
        print(f"  ERROR PIB por el gasto / poblacion: {exc}")
    return ok


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--local", type=Path)
    a = ap.parse_args()
    sys.exit(0 if actualizar(a.local) else 1)
